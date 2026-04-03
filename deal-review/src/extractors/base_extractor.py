import base64
import io
import json
import logging
import os
from abc import ABC, abstractmethod

import anthropic
from tenacity import retry, stop_after_attempt, wait_exponential

logger = logging.getLogger(__name__)

MODEL = os.getenv("CLAUDE_MODEL", "claude-opus-4-6")
MAX_TOKENS = int(os.getenv("CLAUDE_MAX_TOKENS_EXTRACTION", "4096"))
MAX_EXCEL_CHARS = int(os.getenv("CLAUDE_MAX_EXCEL_CHARS", "50000"))
MAX_PDF_BYTES = int(os.getenv("CLAUDE_MAX_PDF_BYTES", str(10 * 1024 * 1024)))


class BaseExtractor(ABC):

    def __init__(self):
        self.client = anthropic.Anthropic(api_key=os.getenv("ANTHROPIC_API_KEY"))

    @abstractmethod
    def build_prompt(self) -> str:
        """Return the extraction prompt for this document type."""

    # ------------------------------------------------------------------
    # Public extraction methods
    # ------------------------------------------------------------------

    def extract_from_pdf(self, pdf_bytes: bytes) -> dict:
        if len(pdf_bytes) > MAX_PDF_BYTES:
            logger.info("PDF too large for document API, falling back to text extraction")
            return self._extract_from_pdf_text(pdf_bytes)

        b64 = base64.standard_b64encode(pdf_bytes).decode("utf-8")
        messages = [
            {
                "role": "user",
                "content": [
                    {
                        "type": "document",
                        "source": {
                            "type": "base64",
                            "media_type": "application/pdf",
                            "data": b64,
                        },
                    },
                    {"type": "text", "text": self.build_prompt()},
                ],
            }
        ]
        return self._call_claude(messages)

    def extract_from_image(self, image_bytes: bytes, media_type: str) -> dict:
        b64 = base64.standard_b64encode(image_bytes).decode("utf-8")
        messages = [
            {
                "role": "user",
                "content": [
                    {
                        "type": "image",
                        "source": {
                            "type": "base64",
                            "media_type": media_type,
                            "data": b64,
                        },
                    },
                    {"type": "text", "text": self.build_prompt()},
                ],
            }
        ]
        return self._call_claude(messages)

    def extract_from_excel(self, excel_bytes: bytes) -> dict:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
            text = self._excel_to_text(wb)
        except Exception as e:
            logger.warning(f"openpyxl failed: {e}. Attempting xlrd fallback.")
            text = self._excel_to_text_xlrd(excel_bytes)

        prompt = self.build_prompt() + "\n\nSPREADSHEET CONTENT:\n" + text
        messages = [{"role": "user", "content": prompt}]
        return self._call_claude(messages)

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------

    def _extract_from_pdf_text(self, pdf_bytes: bytes) -> dict:
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
            text = "\n".join(
                page.extract_text() or "" for page in reader.pages
            )[:MAX_EXCEL_CHARS]
        except Exception as e:
            logger.error(f"pypdf text extraction failed: {e}")
            return {"_parse_error": str(e)}

        prompt = self.build_prompt() + "\n\nDOCUMENT TEXT:\n" + text
        messages = [{"role": "user", "content": prompt}]
        return self._call_claude(messages)

    def _excel_to_text(self, workbook) -> str:
        lines = []
        for sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
            lines.append(f"\n=== SHEET: {sheet_name} ===")
            for row in ws.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    lines.append(
                        "\t".join(str(c) if c is not None else "" for c in row)
                    )
        return "\n".join(lines)[:MAX_EXCEL_CHARS]

    def _excel_to_text_xlrd(self, excel_bytes: bytes) -> str:
        import xlrd
        wb = xlrd.open_workbook(file_contents=excel_bytes)
        lines = []
        for sheet in wb.sheets():
            lines.append(f"\n=== SHEET: {sheet.name} ===")
            for row_idx in range(sheet.nrows):
                row = [str(sheet.cell_value(row_idx, c)) for c in range(sheet.ncols)]
                if any(v.strip() for v in row):
                    lines.append("\t".join(row))
        return "\n".join(lines)[:MAX_EXCEL_CHARS]

    @retry(stop=stop_after_attempt(3), wait=wait_exponential(min=2, max=10))
    def _call_claude(self, messages: list) -> dict:
        response = self.client.messages.create(
            model=MODEL,
            max_tokens=MAX_TOKENS,
            messages=messages,
        )
        return self._parse_response(response.content[0].text)

    def _parse_response(self, text: str) -> dict:
        text = text.strip()
        # Strip markdown code fences if present
        if text.startswith("```"):
            parts = text.split("```")
            if len(parts) >= 3:
                text = parts[1]
            elif len(parts) >= 2:
                text = parts[1]
            if text.startswith("json"):
                text = text[4:]
            text = text.strip()

        try:
            return json.loads(text)
        except json.JSONDecodeError:
            logger.warning(f"JSON parse failed, raw response: {text[:300]}")
            return {"_parse_error": text[:500]}
