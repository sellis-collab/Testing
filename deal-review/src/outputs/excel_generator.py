import logging
import os
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models.deal import DealRecord

logger = logging.getLogger(__name__)

HEADERS = [
    "Deal ID", "Received Date", "Property Name", "Market", "Submarket",
    "Size (SF)", "Clear Height (ft)", "Land (Acres)", "# Buildings",
    "Developer / Sponsor", "Tenant Type", "Pre-Lease Status",
    "Development Cost ($)", "Cost / SF ($)", "Investment Amount ($)", "LTC / LTV (%)",
    "Rent NNN / SF ($)", "Yield-on-Cost (%)", "Market Cap Rate (%)",
    "Spread (bps)", "IRR Levered (%)", "IRR Unlevered (%)",
    "Equity Multiple (x)", "Hold Period (yrs)",
    "GS Rent Growth (%)", "GS Vacancy (%)", "GS Market Score",
    "Score (%)", "Recommendation", "Flags",
]

HEADER_FILL = PatternFill("solid", fgColor="1A3A5C")
HEADER_FONT = Font(bold=True, color="FFFFFF")
PURSUE_FILL = PatternFill("solid", fgColor="C6EFCE")
WATCH_FILL = PatternFill("solid", fgColor="FFEB9C")
PASS_FILL = PatternFill("solid", fgColor="FFC7CE")

DEFAULT_OUTPUT_PATH = "data/output/deal_log.xlsx"


class ExcelGenerator:

    def __init__(self, output_path: str = DEFAULT_OUTPUT_PATH):
        self.output_path = output_path

    def generate(self, deal: DealRecord) -> str:
        Path(self.output_path).parent.mkdir(parents=True, exist_ok=True)

        if os.path.exists(self.output_path):
            wb = openpyxl.load_workbook(self.output_path)
            ws = wb.active
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Deal Log"
            self._write_header(ws)

        ws.append(self._deal_to_row(deal))

        # Colour-code Recommendation cell (column 29 = index of "Recommendation")
        rec_col = HEADERS.index("Recommendation") + 1
        rec_cell = ws.cell(row=ws.max_row, column=rec_col)
        if deal.recommendation == "Pursue":
            rec_cell.fill = PURSUE_FILL
        elif deal.recommendation == "Watch":
            rec_cell.fill = WATCH_FILL
        else:
            rec_cell.fill = PASS_FILL

        wb.save(self.output_path)
        logger.info(f"Excel deal log updated: {self.output_path}")
        return self.output_path

    def _write_header(self, ws):
        ws.append(HEADERS)
        for col_idx, _ in enumerate(HEADERS, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(col_idx)].width = 18
        ws.freeze_panes = "A2"

    def _deal_to_row(self, deal: DealRecord) -> list:
        def pct(v):
            return round(v * 100, 2) if v is not None else None

        return [
            deal.deal_id,
            str(deal.received_date) if deal.received_date else "",
            deal.property_name,
            deal.market,
            deal.submarket,
            deal.size_sf,
            deal.clear_height_ft,
            deal.land_acres,
            deal.num_buildings,
            deal.developer_sponsor,
            deal.tenant_type,
            deal.pre_lease_status,
            deal.development_cost,
            deal.cost_per_sf,
            deal.investment_amount,
            pct(deal.ltc_ltv),
            deal.projected_rent_nnn_sf,
            pct(deal.yield_on_cost),
            pct(deal.market_cap_rate),
            deal.spread_bps,
            pct(deal.irr_levered),
            pct(deal.irr_unlevered),
            deal.equity_multiple,
            deal.hold_period_years,
            pct(deal.gs_market_rent_growth),
            pct(deal.gs_vacancy_rate),
            deal.gs_market_score,
            deal.score_total,
            deal.recommendation,
            " | ".join(deal.score_flags),
        ]
